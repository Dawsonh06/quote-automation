import requests
import json
import os

GRAPH_BASE = "https://graph.microsoft.com/v1.0"

def get_site_id(token, sharepoint_site):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"{GRAPH_BASE}/sites/{sharepoint_site}"
    response = requests.get(url, headers=headers)
    response.raise_for_status()
    return response.json()["id"]

def get_drive_id(token, site_id):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"{GRAPH_BASE}/sites/{site_id}/drives"
    response = requests.get(url, headers=headers)
    response.raise_for_status()
    drives = response.json().get("value", [])
    for drive in drives:
        if drive.get("name") == "Documents":
            return drive["id"]
    return drives[0]["id"]

def upload_file(token, drive_id, folder_path, filename, file_bytes):
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/octet-stream"}
    path = f"{folder_path}/{filename}".replace("//", "/")
    url = f"{GRAPH_BASE}/drives/{drive_id}/root:/{path}:/content"
    response = requests.put(url, headers=headers, data=file_bytes)
    response.raise_for_status()
    return response.json().get("webUrl")

def read_json_file(token, drive_id, file_path):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"{GRAPH_BASE}/drives/{drive_id}/root:/{file_path}:/content"
    response = requests.get(url, headers=headers)
    if response.status_code == 404:
        return {}
    response.raise_for_status()
    return response.json()

def write_json_file(token, drive_id, file_path, data):
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    url = f"{GRAPH_BASE}/drives/{drive_id}/root:/{file_path}:/content"
    response = requests.put(url, headers=headers, data=json.dumps(data, indent=2))
    response.raise_for_status()

def ensure_folder_exists(token, drive_id, folder_path):
    """Create folder_path (relative to drive root) if it does not already exist.

    Walks each path component in order so nested paths are created correctly.
    Raises RuntimeError if any component cannot be created.
    """
    auth_header = {"Authorization": f"Bearer {token}"}
    json_headers = {**auth_header, "Content-Type": "application/json"}
    parts = [p for p in folder_path.replace("\\", "/").split("/") if p]
    current_path = ""

    for part in parts:
        parent_url = (
            f"{GRAPH_BASE}/drives/{drive_id}/root/children"
            if not current_path
            else f"{GRAPH_BASE}/drives/{drive_id}/root:/{current_path}:/children"
        )
        current_path = f"{current_path}/{part}" if current_path else part

        check = requests.get(
            f"{GRAPH_BASE}/drives/{drive_id}/root:/{current_path}",
            headers=auth_header,
        )
        if check.status_code == 200:
            continue

        resp = requests.post(
            parent_url,
            headers=json_headers,
            json={"name": part, "folder": {}, "@microsoft.graph.conflictBehavior": "ignore"},
        )
        if resp.status_code not in (200, 201):
            raise RuntimeError(
                f"Could not create folder '{current_path}': "
                f"{resp.status_code} {resp.text[:200]}"
            )


def append_to_log(token, drive_id, log_path, row_data):
    import openpyxl
    import io
    from datetime import datetime

    headers_auth = {"Authorization": f"Bearer {token}"}
    url = f"{GRAPH_BASE}/drives/{drive_id}/root:/{log_path}:/content"
    response = requests.get(url, headers=headers_auth)

    if response.status_code == 404:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Processing Log"
        ws.append(["Timestamp", "Subject", "Sender", "Classification",
                   "Matched Bid", "Confidence", "Vendor", "Scope",
                   "Price", "Filename", "Filed To", "Notes"])
    else:
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active

    from zoneinfo import ZoneInfo
    _pst = datetime.now(ZoneInfo("America/Los_Angeles")).strftime("%Y-%m-%d %H:%M PST")
    ws.insert_rows(2)
    for col_idx, value in enumerate([_pst] + row_data, start=1):
        ws.cell(row=2, column=col_idx, value=value)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    requests.put(url, headers={**headers_auth, "Content-Type": "application/octet-stream"},
                 data=buffer.read()).raise_for_status()
