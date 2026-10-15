import os
import io
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
from dotenv import load_dotenv

load_dotenv()

GRAPH_BASE = "https://graph.microsoft.com/v1.0"
DB_PATH = "Subcontractor Database.xlsx"

HEADERS = [
    "Company Name", "Email", "Phone", "Location",
    "Categories", "Bids Quoted", "Last Quote Date", "Total Quotes"
]

def load_database(token, drive_id):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"{GRAPH_BASE}/drives/{drive_id}/root:/{DB_PATH}:/content"
    response = requests.get(url, headers=headers)

    if response.status_code == 404:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Subcontractors"
        # Write headers
        header_fill = PatternFill("solid", start_color="1F4E78")
        header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        for col, h in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        # Column widths
        widths = [30, 35, 18, 25, 40, 60, 18, 14]
        from openpyxl.utils import get_column_letter
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return wb

    return openpyxl.load_workbook(io.BytesIO(response.content))


def save_database(token, drive_id, wb):
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/octet-stream"}
    url = f"{GRAPH_BASE}/drives/{drive_id}/root:/{DB_PATH}:/content"
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    requests.put(url, headers=headers, data=buffer.read()).raise_for_status()


def upsert_vendor(token, drive_id, vendor_info):
    """
    vendor_info keys: company, email, phone, location, category_short, bid_name, quote_date
    Adds new vendor or updates existing one.
    """
    company = str(vendor_info.get("company") or "Unknown").strip()
    email = str(vendor_info.get("email") or "").strip()
    phone = str(vendor_info.get("phone") or "").strip()
    location = str(vendor_info.get("location") or "").strip()
    category = str(vendor_info.get("category_short") or "").strip()
    bid_name = str(vendor_info.get("bid_name") or "").strip()
    quote_date = vendor_info.get("quote_date") or datetime.today().strftime("%m/%d/%Y")

    wb = load_database(token, drive_id)
    ws = wb.active

    # Find existing row for this vendor
    existing_row = None
    for row in ws.iter_rows(min_row=2):
        if row[0].value and str(row[0].value).strip().lower() == company.lower():
            existing_row = row
            break

    if existing_row:
        # Update existing vendor
        # Add category if not already listed
        existing_cats = str(existing_row[4].value or "")
        if category and category not in existing_cats:
            existing_row[4].value = f"{existing_cats}; {category}".strip("; ")

        # Append bid if not already listed
        existing_bids = str(existing_row[5].value or "")
        if bid_name and bid_name not in existing_bids:
            existing_row[5].value = f"{existing_bids}; {bid_name}".strip("; ")

        # Update last quote date and total
        existing_row[6].value = quote_date
        existing_row[7].value = (existing_row[7].value or 0) + 1

        # Fill in blanks if we have better data now
        if not existing_row[1].value and email:
            existing_row[1].value = email
        if not existing_row[2].value and phone:
            existing_row[2].value = phone
        if not existing_row[3].value and location:
            existing_row[3].value = location

        print(f"  Updated vendor: {company}")
    else:
        # Add new vendor row
        row_font = Font(name="Calibri", size=11)
        new_row = [company, email, phone, location, category, bid_name, quote_date, 1]
        ws.append(new_row)
        for cell in ws[ws.max_row]:
            cell.font = row_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        print(f"  Added new vendor: {company}")

    save_database(token, drive_id, wb)
